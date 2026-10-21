#!/usr/bin/env python3
"""Read the supplied clinical workbook and join only explicitly crosswalked people.

Source diagnosis labels describe availability; they do not define a binary outcome.
The user confirmed that this workbook's age is age at sampling; missing ages stay missing.
"""

import argparse
import hashlib
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from svgp_inventory import SAMPLE_FIELDS, known, protect_inputs, read_table, unique, write_table


SOURCE_FIELDS = ["SequencingID_final", "Sample_ID_WGS_given", "Sample_ID_Biobank_raw",
                 "sex", "age", "Diagnosis_svPPA", "Diagnosis_complete", "DNA_from"]
PHENOTYPE_FIELDS = ["person_id", "workbook_sample_id", "workbook_sequencing_id",
                    "workbook_biobank_id", "phenotype", "source_phenotype_column",
                    "diagnosis_svppa", "diagnosis_complete", "source_sex", "source_age",
                    "age_at_sampling", "age_definition",
                    "dna_from", "source_workbook", "source_sheet", "source_row",
                    "source_workbook_sha256", "join_status"]


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sampling_age(value, label):
    if not known(value):
        return "unknown"
    try:
        age = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label}: age at sampling must be a non-negative finite number: {value!r}") from None
    if not math.isfinite(age) or age < 0:
        raise ValueError(f"{label}: age at sampling must be a non-negative finite number: {value!r}")
    return str(int(age)) if age.is_integer() else str(age)


def import_workbook(args):
    output = Path(args.output)
    names = ["phenotypes.tsv", "phenotype_join_report.tsv", "phenotype_import_summary.tsv"]
    if args.samples:
        names.append("samples_with_phenotypes.tsv")
    protect_inputs([args.workbook, args.crosswalk, args.samples], [output / name for name in names])
    with Path(args.workbook).open("rb") as handle:
        checksum = hashlib.file_digest(handle, "sha256").hexdigest()
    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    try:
        if args.sheet not in workbook.sheetnames:
            raise ValueError(f"Workbook has no sheet {args.sheet!r}")
        sheet_rows = workbook[args.sheet].iter_rows()
        headers = [cell_text(cell.value) for cell in next(sheet_rows, [])]
        required = set(SOURCE_FIELDS + [args.id_column, args.phenotype_column])
        if any(not header for header in headers) or len(headers) != len(set(headers)) or not required.issubset(headers):
            raise ValueError("Workbook has duplicate headers or is missing required source columns")
        source = []
        for number, cells in enumerate(sheet_rows, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise ValueError(f"Workbook row {number}: formula/error cells require reviewed values")
            row = dict(zip(headers, [cell_text(cell.value) for cell in cells]))
            row["source_row"] = number
            source.append(row)
    finally:
        workbook.close()
    source_by_id = unique(source, args.id_column, "phenotype workbook")
    mappings = read_table(args.crosswalk, [args.id_column, "person_id"])
    unique(mappings, args.id_column, "phenotype crosswalk")
    unique(mappings, "person_id", "phenotype crosswalk; choose one workbook row per person")
    samples = read_table(args.samples, SAMPLE_FIELDS) if args.samples else []
    people = unique(samples, "person_id", "curated samples")
    person_by_source = {}
    for row in mappings:
        source_id, person = row[args.id_column], row["person_id"]
        if source_id not in source_by_id:
            raise ValueError(f"Crosswalk source ID is absent from workbook: {source_id!r}")
        if args.samples and person not in people:
            raise ValueError(f"Crosswalk person is absent from curated samples: {person!r}")
        person_by_source[source_id] = person

    records = []
    for row in source:
        source_id = row[args.id_column]
        person = person_by_source.get(source_id, "unknown")
        record = dict(person_id=person, workbook_sample_id=source_id,
                      workbook_sequencing_id=row["SequencingID_final"],
                      workbook_biobank_id=row["Sample_ID_Biobank_raw"],
                      phenotype=row[args.phenotype_column] or "unknown",
                      source_phenotype_column=args.phenotype_column,
                      diagnosis_svppa=row["Diagnosis_svPPA"], diagnosis_complete=row["Diagnosis_complete"],
                      source_sex=row["sex"], source_age=row["age"],
                      age_at_sampling=sampling_age(row["age"], f"Workbook row {row['source_row']}"),
                      age_definition="age_at_sampling",
                      dna_from=row["DNA_from"] or "unknown", source_workbook=args.workbook,
                      source_sheet=args.sheet, source_row=row["source_row"], source_workbook_sha256=checksum,
                      join_status="MAPPED" if known(person) else "UNMAPPED")
        records.append(record)
        if person in people:
            sample = people[person]
            for field, value in (("phenotype", record["phenotype"]), ("sex", record["source_sex"])):
                if known(sample[field]) and known(value) and sample[field] != value:
                    raise ValueError(f"{person}: curated {field} conflicts with workbook row {row['source_row']}")
                if not known(sample[field]) and known(value):
                    sample[field] = value
            # The explicit crosswalk can connect an existing, differently named phenotype ID.
            if not known(sample["phenotype_id"]):
                sample["phenotype_id"] = source_id
            if known(sample["age_definition"]) and sample["age_definition"] != "age_at_sampling":
                raise ValueError(f"{person}: curated age_definition conflicts with workbook age at sampling")
            curated_age = sampling_age(sample.get("age_at_sampling", "unknown"), f"{person}: curated")
            source_age = record["age_at_sampling"]
            if known(curated_age) and known(source_age) and float(curated_age) != float(source_age):
                raise ValueError(f"{person}: curated age_at_sampling conflicts with workbook row {row['source_row']}")
            if not known(curated_age):
                sample["age_at_sampling"] = source_age
            sample["age_definition"] = "age_at_sampling"
            sample.update({key: value for key, value in record.items()
                           if key not in {"person_id", "phenotype", "age_at_sampling", "age_definition"}})

    write_table(output / "phenotypes.tsv", records, PHENOTYPE_FIELDS)
    report_fields = ["workbook_sample_id", "person_id", "source_row", "join_status"]
    write_table(output / "phenotype_join_report.tsv", records, report_fields)
    summary = [dict(metric="workbook_rows", category="all", count=len(records)),
               dict(metric="curated_people", category="all", count=len(samples))]
    for metric, values in (("join_status", [row["join_status"] for row in records]),
                           ("source_phenotype", [row["phenotype"] for row in records])):
        summary.extend(dict(metric=metric, category=category, count=count)
                       for category, count in sorted(Counter(values).items()))
    write_table(output / "phenotype_import_summary.tsv", summary, ["metric", "category", "count"])
    if args.samples:
        with Path(args.samples).open() as handle:
            fields = handle.readline().rstrip("\r\n").split("\t")
        fields += [field for field in PHENOTYPE_FIELDS if field not in fields]
        write_table(output / "samples_with_phenotypes.tsv", samples, fields)
    print(f"Read {len(records)} clinical rows; explicitly linked {len(mappings)} people.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", default="Blad1")
    parser.add_argument("--id-column", default="Sample_ID_WGS_given")
    parser.add_argument("--phenotype-column", default="Diagnosis_svPPA")
    parser.add_argument("--crosswalk", required=True)
    parser.add_argument("--samples", help="Existing one-person-per-row curated metadata TSV")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        import_workbook(args)
    except (OSError, ValueError) as error:
        parser.exit(1, f"Phenotype import failed: {error}\n")


if __name__ == "__main__":
    main()
